"""Japan MHLW - violation cases of imported foods (monthly / cumulative Excel, Japanese)."""
import re, io, hashlib, datetime as dt
from html.parser import HTMLParser
from common import http, http_text, clean, is_india, since_date

INDEX = 'https://www.mhlw.go.jp/stf/seisakunitsuite/bunya/kenkou_iryou/shokuhin/yunyu_kanshi/ihan/index.html'
BASE = 'https://www.mhlw.go.jp'


class Anchors(HTMLParser):
    def __init__(self):
        super().__init__(); self.items = []; self._h = None; self._t = []
    def handle_starttag(self, tag, attrs):
        if tag == 'a':
            h = dict(attrs).get('href') or ''
            if re.search(r'\.xlsx?$', h, re.I): self._h, self._t = h, []
    def handle_data(self, d):
        if self._h is not None: self._t.append(d)
    def handle_endtag(self, tag):
        if tag == 'a' and self._h is not None:
            self.items.append((self._h, clean(''.join(self._t)))); self._h = None


def jp_date(v):
    if v is None or v == '': return ''
    if isinstance(v, (dt.datetime, dt.date)):
        return v.date().isoformat() if isinstance(v, dt.datetime) else v.isoformat()
    if isinstance(v, (int, float)) and 20000 < v < 80000:
        return (dt.date(1899, 12, 30) + dt.timedelta(days=int(v))).isoformat()
    s = str(v).strip()
    m = re.search(r'(令和|平成|R|H)\s*(\d{1,2}|元)\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日', s)
    if m:
        y = 1 if m[2] == '元' else int(m[2]); base = 2018 if m[1] in ('令和', 'R') else 1988
        return f'{base + y:04d}-{int(m[3]):02d}-{int(m[4]):02d}'
    m = re.search(r'(\d{4})[年/.\-](\d{1,2})[月/.\-](\d{1,2})', s)
    if m: return f'{int(m[1]):04d}-{int(m[2]):02d}-{int(m[3]):02d}'
    return ''


def sheets(raw, name):
    """Yield list-of-rows for every sheet of an .xls/.xlsx."""
    if name.lower().endswith('.xlsx'):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
        for ws in wb.worksheets:
            yield [list(r) for r in ws.iter_rows(values_only=True)]
    else:
        import xlrd
        wb = xlrd.open_workbook(file_contents=raw)
        for sh in wb.sheets():
            rows = []
            for i in range(sh.nrows):
                r = []
                for j in range(sh.ncols):
                    c = sh.cell(i, j)
                    v = c.value
                    if c.ctype == xlrd.XL_CELL_DATE:
                        v = xlrd.xldate_as_datetime(v, wb.datemode)
                    r.append(v)
                rows.append(r)
            yield rows


COLS = {'date': r'年月日|通知|処分|日付', 'product': r'品名|食品等', 'origin': r'原産国|輸出国|国名|生産国', 'law': r'違反条文|条文', 'content': r'違反内容|内容|理由', 'action': r'措置|処置'}


def parse_rows(rows):
    """Find the header row, map columns by Japanese keywords, return India rows."""
    hi = None
    for i, r in enumerate(rows[:40]):
        txt = [str(c or '') for c in r]
        if any(re.search(r'品名', t) for t in txt) and any(re.search(COLS['origin'], t) for t in txt):
            hi = i; break
    if hi is None:
        return None
    head = [str(c or '') for c in rows[hi]]
    idx = {}
    for k, rx in COLS.items():
        for j, h in enumerate(head):
            if re.search(rx, h) and j not in idx.values():
                idx[k] = j; break
    if 'product' not in idx or 'origin' not in idx:
        return None
    out = []
    for r in rows[hi + 1:]:
        if idx['origin'] >= len(r) or not is_india(r[idx['origin']]): continue
        g = lambda k: clean(r[idx[k]]) if k in idx and idx[k] < len(r) else ''
        date = jp_date(r[idx['date']]) if 'date' in idx and idx['date'] < len(r) else ''
        if not date: continue
        product, law, content = g('product'), g('law'), g('content')
        key = '|'.join([date, product, law, content])
        out.append({'id': 'JP-' + hashlib.md5(key.encode()).hexdigest()[:12], 'mk': 'JP', 'date': date,
                    'product': product, 'hazard': f'{law} {content}'.strip(), 'risk': 0, 'ref': '',
                    'desc': clean(f'{product}. {content or law}', 130)})
    return out


def pick_files(items):
    cum = [(h, t) for h, t in items if '輸入' in t and re.search(r'[～〜~]', t)]
    if cum: return cum[:2]
    return [(h, t) for h, t in items if '国内' not in t][:14]


def fetch():
    items = Anchors(); items.feed(http_text(INDEX)); items = items.items
    files = pick_files(items)
    if not files:
        raise RuntimeError('no Excel links found on the MHLW page (page layout may have changed)')
    since = since_date('jp').isoformat()
    recs, preview = [], None
    for href, label in files:
        url = href if href.startswith('http') else BASE + href
        raw = http(url, timeout=120)
        ok = False
        for rows in sheets(raw, href):
            got = parse_rows(rows)
            if got is None:
                preview = preview or rows[:6]
                continue
            ok = True; recs += [r for r in got if r['date'] >= since]
        if not ok:
            raise RuntimeError(f'could not find header columns in {url}; first rows: {preview}')
    return recs
